"""
docx_styler.py - Motor de transformación y formateo editorial de documentos docx para PSLL.
Aplica la identidad corporativa completa de forma automatizada y reproducible.
"""

import os
import re
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from brand_theme import (
    HEX_COLORS, RGB_COLORS, FONT_HEADINGS, FONT_BODY,
    set_cell_shading, set_cell_margins, set_callout_borders, set_table_borders
)

def set_document_geometry(doc):
    """Configura márgenes estándar institucionales de 2.5 cm."""
    for section in doc.sections:
        section.top_margin = Inches(0.98)     # ~2.5 cm
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

def setup_document_styles(doc):
    """Establece las tipografías y colores base en los estilos del documento."""
    styles = doc.styles
    
    # Normal (Cuerpo de texto)
    if 'Normal' in styles:
        normal = styles['Normal']
        normal.font.name = FONT_BODY
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGB_COLORS['ink_green']
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(4.5)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Heading 1 (Capítulo / Bloque mayor)
    if 'Heading 1' in styles:
        h1 = styles['Heading 1']
        h1.font.name = FONT_HEADINGS
        h1.font.size = Pt(15.5)
        h1.font.bold = True
        h1.font.color.rgb = RGB_COLORS['deep_green']
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(4)
        h1.paragraph_format.keep_with_next = True
        
    # Heading 2 (Subepígrafe nivel 2)
    if 'Heading 2' in styles:
        h2 = styles['Heading 2']
        h2.font.name = FONT_HEADINGS
        h2.font.size = Pt(12.5)
        h2.font.bold = True
        h2.font.color.rgb = RGB_COLORS['sage']
        h2.paragraph_format.space_before = Pt(10)
        h2.paragraph_format.space_after = Pt(3)
        h2.paragraph_format.keep_with_next = True

    # Heading 3 (Subepígrafe nivel 3)
    if 'Heading 3' in styles:
        h3 = styles['Heading 3']
        h3.font.name = FONT_HEADINGS
        h3.font.size = Pt(11.0)
        h3.font.bold = True
        h3.font.color.rgb = RGB_COLORS['deep_green']
        h3.paragraph_format.space_before = Pt(7)
        h3.paragraph_format.space_after = Pt(2)
        h3.paragraph_format.keep_with_next = True

def add_header_and_footer(doc, topic_title: str):
    """Agrega cabecera y pie de página dinámico con paginación en Word."""
    for section in doc.sections:
        # Encabezado
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.text = f"POLÍTICAS SOCIOLABORALES Y DE EMPLEO · {topic_title.upper()}"
        hp.runs[0].font.name = FONT_HEADINGS
        hp.runs[0].font.size = Pt(8.0)
        hp.runs[0].font.color.rgb = RGB_COLORS['sage']
        
        # Pie de página
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Texto izquierdo en pie
        run_l = fp.add_run("Universidad Pablo de Olavide (UPO) · Grado en RRLL y RRHH | ")
        run_l.font.name = FONT_BODY
        run_l.font.size = Pt(8.5)
        run_l.font.color.rgb = RGB_COLORS['sage']
        
        # Paginación dinámica XML: Página X de Y
        run_p = fp.add_run("Página ")
        run_p.font.name = FONT_BODY
        run_p.font.size = Pt(8.5)
        run_p.font.color.rgb = RGB_COLORS['ink_green']
        
        fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fld_page)
        
        run_de = fp.add_run(" de ")
        run_de.font.name = FONT_BODY
        run_de.font.size = Pt(8.5)
        run_de.font.color.rgb = RGB_COLORS['ink_green']
        
        fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        fp._p.append(fld_numpages)

def insert_callout_box(doc, callout_type: str, title: str, text: str):
    """
    Inserta una caja destacada (callout box) de 1 celda con sombreado y borde izquierdo grueso.
    Tipos: 'session', 'concept', 'warning', 'case'.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    if callout_type == "session":
        bg_hex = HEX_COLORS["mint"]          # Verde menta suave
        border_hex = HEX_COLORS["deep_green"] # Borde verde profundo distintivo
        icon = "🗓️"
        title_color = RGB_COLORS["deep_green"]
    elif callout_type == "concept":
        bg_hex = HEX_COLORS["mint"]
        border_hex = HEX_COLORS["sage"]
        icon = "💡"
        title_color = RGB_COLORS["deep_green"]
    elif callout_type == "warning":
        bg_hex = HEX_COLORS["alert_bg"]
        border_hex = HEX_COLORS["coral"]
        icon = "⚠️"
        title_color = RGB_COLORS["coral"]
    else: # case / empirical
        bg_hex = HEX_COLORS["warm_cream"]
        border_hex = HEX_COLORS["deep_green"]
        icon = "📊"
        title_color = RGB_COLORS["deep_green"]
        
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top_dpt=140, bottom_dpt=140, left_dpt=200, right_dpt=160)
    set_callout_borders(cell, border_hex, border_size_pt=26)
    
    # Párrafo del título
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_icon = p_title.add_run(f"{icon}  {title.upper()}\n")
    run_icon.font.name = FONT_HEADINGS
    run_icon.font.size = Pt(9.5)
    run_icon.font.bold = True
    run_icon.font.color.rgb = title_color
    
    # Texto interior (justificado)
    p_text = cell.add_paragraph()
    p_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_text.paragraph_format.space_before = Pt(0)
    p_text.paragraph_format.space_after = Pt(0)
    p_text.paragraph_format.line_spacing = 1.15
    run_text = p_text.add_run(text)
    run_text.font.name = FONT_BODY
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGB_COLORS["ink_green"]
    
    # Espacio tras la caja
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)


def insert_figure(doc, image_path: str, caption_text: str, source_text: str, width_inches=6.2):
    """Inserta una figura centrada con pie y fuente reglamentarios."""
    if not os.path.exists(image_path):
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))
    
    # Pie de figura
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(1)
    p_cap.paragraph_format.keep_with_next = True
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = FONT_HEADINGS
    run_cap.font.size = Pt(9.0)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGB_COLORS["deep_green"]
    
    # Fuente
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_before = Pt(0)
    p_src.paragraph_format.space_after = Pt(8)
    run_src = p_src.add_run(source_text)
    run_src.font.name = FONT_BODY
    run_src.font.size = Pt(8.0)
    run_src.font.italic = True
    run_src.font.color.rgb = RGBColor(111, 128, 115)
