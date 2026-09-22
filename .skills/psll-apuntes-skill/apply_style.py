"""
apply_style.py - Script principal de la skill psll-apuntes-skill.
Lee un documento de apuntes en .docx, aplica la maquetación corporativa completa,
inserta figuras de alta resolución, cajas callout y tablas estilizadas, y guarda el resultado.
"""

import os
import sys
import re
import io
import argparse
import docx

if hasattr(sys.stdout, 'buffer'):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
if hasattr(sys.stderr, 'buffer'):
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')
import zipfile
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Asegurar importación de módulos hermanos
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
if CURRENT_DIR not in sys.path:
    sys.path.insert(0, CURRENT_DIR)

from brand_theme import (
    HEX_COLORS, RGB_COLORS, FONT_HEADINGS, FONT_BODY,
    set_cell_shading, set_cell_margins, set_callout_borders, set_table_borders
)
from docx_styler import (
    set_document_geometry, setup_document_styles, add_header_and_footer,
    insert_callout_box, insert_reflection_box, insert_key_concepts_box, insert_figure
)

TEMA_1_FIGURES_META = {
    "image1.png": {
        "remastered": "epa_taxonomy.png",
        "caption": "Figura 1.1: Taxonomía oficial y articulación de la población según la EPA (INE / OIT).",
        "source": "Fuente: Elaboración propia a partir de la metodología de la Encuesta de Población Activa (INE)."
    },
    "image2.png": {
        "remastered": "epa_decision_tree.png",
        "caption": "Figura 1.2: Algoritmo y criterios de clasificación de la población laboral según la OIT y la EPA.",
        "source": "Fuente: Elaboración propia a partir de los criterios metodológicos de la OIT y la Encuesta de Población Activa (INE)."
    },
    "image3.png": {
        "remastered": "vab_pan.png",
        "caption": "Figura 1.3: Cadena de valor añadido bruto (VAB) y proceso de producción del pan.",
        "source": "Fuente: Elaboración propia para Políticas Sociolaborales (UPO). Metodología de Contabilidad Nacional (INE / SEC-2010)."
    },
    "image4.jpeg": {
        "remastered": "pib_interanual.png",
        "caption": "Figura 1.4: Evolución del Producto Interior Bruto (PIB) en España: tasas de variación interanual en volumen encadenado (2022T1 - 2026T2).",
        "source": "Fuente: Instituto Nacional de Estadística (INE), Contabilidad Nacional Trimestral de España (CNTR)."
    },
    "image5.png": {
        "remastered": "productividad_salarios.png",
        "caption": "Figura 1.5: Relación entre productividad por hora trabajada y salario medio en la OCDE (2024).",
        "source": "Fuente: Elaboración propia a partir de datos oficiales de la OCDE (Productivity Statistics & Employment Outlook, 2024)."
    },
    "image6.jpeg": {
        "remastered": "modelo_desempleo_neoclasico.png",
        "caption": "Figura 1.6: El desempleo en el modelo neoclásico del mercado de trabajo: rigidez salarial y exceso de oferta.",
        "source": "Fuente: Elaboración propia a partir del modelo neoclásico del mercado de trabajo (adaptado de M. Barneto)."
    },
    "image7.png": {
        "remastered": "tasa_natural_nairu.png",
        "caption": "Figura 1.7: Tasa de desempleo observada y tasa natural de desempleo (NAIRU) en España y la Zona Euro (1980–2024): desempleo estructural vs. desempleo cíclico.",
        "source": "Fuente: Elaboración propia a partir de datos oficiales de la base AMECO (Comisión Europea, DG ECFIN, 2024)."
    },
    "image8.png": {
        "remastered": "curva_phillips_dual.png",
        "caption": "Figura 1.8: La Curva de Phillips: modelo teórico de expectativas aceleracionistas (Friedman-Phelps) y evidencia empírica en España (2002–2024).",
        "source": "Fuente: Elaboración propia para Políticas Sociolaborales (UPO) a partir de Friedman (1968), Phelps (1967) e Instituto Nacional de Estadística (EPA y ETCL, 2024)."
    },
    "image9.png": {
        "remastered": "curva_beveridge_dual.png",
        "caption": "Figura 1.9: La Curva de Beveridge: modelo teórico de emparejamiento (DMP) y evidencia empírica en España (1980–2024).",
        "source": "Fuente: Elaboración propia para Políticas Sociolaborales (UPO) a partir de FEDEA (2017), Nada es Gratis e Instituto Nacional de Estadística (EPA y ETCL, 2024)."
    },
    "image10.png": {
        "remastered": None,
        "skip": True
    },
    "image11.png": {
        "remastered": None,
        "skip": True
    },
}

TEMA_1_SESSION_BOXES = [
    {
        "id": "sesion_1",
        "pattern": r"^1\.\s+FUNDAMENTOS DEL MERCADO DE TRABAJO",
        "title": "Encuadre Docente Presencial · Sesión 1 (Semana 1 · Lunes) — Marco Analítico, Población y Flujos Laborales",
        "text": "Desarrollo presencial de los conceptos fundamentales del mercado de trabajo: stocks poblacionales (activos, ocupados, parados, inactivos), tasas básicas EPA, matriz de flujos laborales y dinámicas de actividad y desánimo. Apertura con el enigma laboral de inicio de curso, sondeo diagnóstico interactivo y taller de análisis empírico con microdatos oficiales."
    },
    {
        "id": "sesion_2",
        "pattern": r"^2\.\s+EL PIB Y LA PRODUCTIVIDAD EN EL MERCADO LABORAL",
        "title": "Encuadre Docente Presencial · Sesión 2 (Semana 1 · Martes/Viernes) — Macroeconomía del Trabajo, Productividad y CLU",
        "text": "Articulación macroeconómica de la demanda de trabajo: distinción rigurosa entre PIB nominal y PIB real, deflactores, descomposición del PIB per cápita, productividad aparente del trabajo y Costes Laborales Unitarios (CLU). Incluye simulación interactiva con la Micro-App de deflactores y CLU y resolución del dilema de política salarial."
    },
    {
        "id": "sesion_3",
        "pattern": r"^(\d+\.\s+)?EQUILIBRIO EN EL MERCADO DE TRABAJO",
        "title": "Encuadre Docente Presencial · Sesión 3 (Semana 2 · Lunes) — Microfundamentos de Oferta, Demanda y Salario de Reserva",
        "text": "Modelización microeconómica del mercado laboral: decisiones de oferta de trabajo individual (modelo ocio-consumo, efectos renta y sustitución), salario de reserva, trampas de pobreza e inactividad, y derivación de la demanda de empleo a corto y largo plazo bajo competencia e imperfecciones."
    },
    {
        "id": "sesion_4",
        "pattern": r"^4\.\s+TIPOLOG[IÍ]A DEL DESEMPLEO",
        "title": "Encuadre Docente Presencial · Sesión 4 (Semana 2 · Martes/Viernes) — Tipologías de Paro, NAIRU y Curva de Phillips",
        "text": "Disección analítica del desempleo: friccional, estacional, cíclico y estructural (Epígrafes 4 y 5). Modelo de Curva de Phillips con expectativas (Friedman-Phelps), estimación de la tasa natural / NAIRU y contraste empírico de la desinflación y devaluación salarial en la economía española (2002–2024). Reto manuscrito de cálculo analítico de tipos de paro."
    },
    {
        "id": "sesion_5",
        "pattern": r"^6\.\s+LA CURVA DE BEVERIDGE",
        "title": "Encuadre Docente Presencial · Sesión 5 (Semana 3 · Cierre Tema 1) — Emparejamiento (DMP), Curva de Beveridge y Políticas Activas",
        "text": "Modelo de búsqueda y emparejamiento (Diamond-Mortensen-Pissarides): función de emparejamiento, tasa de paro de estado estacionario u* = s/(s+f), tensión del mercado (labor market tightness θ) y análisis comparado de la Curva de Beveridge en España. Síntesis integral de Tema 1 y lanzamiento del Caso 1 de EPD."
    },
]

def extract_media_from_docx(docx_path: str, output_dir: str):
    """
    Extrae automáticamente todas las imágenes del archivo docx original
    a la carpeta relativa figuras/originales/ y devuelve el diccionario rId -> filename.
    """
    os.makedirs(output_dir, exist_ok=True)
    doc_raw = docx.Document(docx_path)
    rid_map = {}
    for rid, rel in doc_raw.part.rels.items():
        if "image" in rel.target_ref:
            rid_map[rid] = os.path.basename(rel.target_ref)
            
    with zipfile.ZipFile(docx_path, 'r') as z:
        for member in z.namelist():
            if member.startswith("word/media/"):
                fname = os.path.basename(member)
                if fname:
                    dest = os.path.join(output_dir, fname)
                    with open(dest, "wb") as f_out:
                        f_out.write(z.read(member))
                        
    print(f"📦 Se han extraído {len(rid_map)} figuras originales a: {output_dir}")
    return rid_map

def create_institutional_cover(doc, emblem_path: str, upo_logo_path: str):
    """Inserta la cabecera institucional en la primera página alineada a la izquierda."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    
    # Celda izquierda: Textos institucionales alineados a la izquierda
    cell_l = tbl.cell(0, 0)
    cell_l.width = Inches(4.8)
    
    # Párrafo 1: UNIVERSIDAD PABLO DE OLAVIDE
    p1 = cell_l.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1.5)
    run_univ = p1.add_run("UNIVERSIDAD PABLO DE OLAVIDE")
    run_univ.font.name = FONT_HEADINGS
    run_univ.font.size = Pt(9.0)
    run_univ.font.bold = True
    run_univ.font.color.rgb = RGB_COLORS["deep_green"]
    
    # Párrafo 2: Facultad de Ciencias del Trabajo...
    p2 = cell_l.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1.5)
    run_fac = p2.add_run("Facultad de Ciencias del Trabajo · Grado en RRLL y Recursos Humanos")
    run_fac.font.name = FONT_BODY
    run_fac.font.size = Pt(8.5)
    run_fac.font.color.rgb = RGB_COLORS["sage"]
    
    # Párrafo 3: Políticas Sociolaborales y de Empleo...
    p3 = cell_l.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(1.5)
    run_asig = p3.add_run("Políticas Sociolaborales y de Empleo (Código 102023) | Curso 2026-2027")
    run_asig.font.name = FONT_BODY
    run_asig.font.size = Pt(8.5)
    run_asig.font.color.rgb = RGB_COLORS["ink_green"]
    
    # Párrafo 4: Prof. Manuel A. Hidalgo Pérez
    p4 = cell_l.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p4.paragraph_format.space_before = Pt(0)
    p4.paragraph_format.space_after = Pt(0)
    run_prof = p4.add_run("Prof. Manuel A. Hidalgo Pérez")
    run_prof.font.name = FONT_BODY
    run_prof.font.size = Pt(8.5)
    run_prof.font.color.rgb = RGB_COLORS["ink_green"]
    
    # Celda derecha: Emblema oficial PSLL
    cell_r = tbl.cell(0, 1)
    cell_r.width = Inches(1.8)
    p_logo = cell_r.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    
    if os.path.exists(emblem_path):
        p_logo.add_run().add_picture(emblem_path, width=Inches(1.8))
        
    # Línea separadora horizontal alineada a la izquierda
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(16)
    run_div = p_div.add_run("―" * 48)
    run_div.font.name = FONT_HEADINGS
    run_div.font.size = Pt(10)
    run_div.font.color.rgb = RGB_COLORS["sage"]

def insert_epa_summary_table(doc):
    """Inserta una tabla estilizada con el resumen de las tasas fundamentales de la EPA."""
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["Indicador Laboral", "Fórmula Matemática Oficial", "Interpretación Económica"]
    widths = [Inches(1.8), Inches(2.2), Inches(2.5)]
    
    # Cabecera
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_shading(hdr_cells[i], HEX_COLORS["deep_green"])
        set_cell_margins(hdr_cells[i], top_dpt=140, bottom_dpt=140, left_dpt=140, right_dpt=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = FONT_HEADINGS
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGB_COLORS["white"]
        
    data = [
        ("Tasa de Actividad", "Activos / Población ≥ 16 × 100", "Mide la propensión y disposición a participar en el mercado de trabajo."),
        ("Tasa de Desempleo (Paro)", "Parados / Población Activa × 100", "Proporción de personas que buscando activamente trabajo no lo encuentran."),
        ("Tasa de Empleo (Ocupación)", "Ocupados / Población ≥ 16 × 100", "Capacidad real de la economía de generar puestos para la población en edad laboral."),
    ]
    
    for row_idx, (col1, col2, col3) in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F7FAF8" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([col1, col2, col3]):
            row_cells[i].width = widths[i]
            set_cell_shading(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top_dpt=120, bottom_dpt=120, left_dpt=140, right_dpt=140)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = FONT_BODY
            run.font.size = Pt(9.0)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = RGB_COLORS["deep_green"]
            elif i == 1:
                run.font.color.rgb = RGB_COLORS["sage"]
                run.font.bold = True
            else:
                run.font.color.rgb = RGB_COLORS["ink_green"]
                
    set_table_borders(table)
    
    # Párrafo posterior
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(2)
    p_post.paragraph_format.space_after = Pt(6)

def extract_paragraph_full_text(p):
    """Extrae el texto completo del párrafo, preservando fórmulas matemáticas de nodos m:oMath."""
    texts = []
    for node in p._p.iter():
        tag = node.tag.split('}')[-1]
        if tag == 't':
            if node.text:
                texts.append(node.text)
    full = ''.join(texts).strip()
    return full if full else p.text.strip()

def clean_and_enhance_text(raw_text: str) -> str:
    """Normaliza fórmulas matemáticas y ajusta redacciones clave a las figuras oficiales."""
    # Fórmulas de la Sección 2 (PIB y Productividad)
    if raw_text == "PIB=VAB+impuestos-subvenciones":
        return "PIB = VAB + Impuestos indirectos netos de subvenciones"
    if "t.c.Xt=Xt-Xt-1Xt-1×100" in raw_text or "t.c.Xt=" in raw_text:
        return "Tasa de variación interanual:  t.c. X_t = [(X_t - X_{t-1}) / X_{t-1}] × 100"
    if raw_text == "PIB per cápita=PIB totalPoblación total":
        return "PIB per cápita = PIB total / Población total"
    if "Productividad Laboral=ProducciónNúmero de trabajadores" in raw_text or "Productividad Laboral=" in raw_text:
        return "Productividad laboral = Producción total (PIB) / (Número de trabajadores o de horas trabajadas)"
    
    # Fórmulas de la Sección 5 (Curva de Phillips)
    if "πt=πte-λut-un" in raw_text or raw_text.startswith("πt="):
        return "π_t = π_t^e - λ · (u_t - u_n)"
    if raw_text == "πt = Tasa de inflación actual" or raw_text.startswith("= Tasa de inflación actual"):
        return "π_t: Tasa de inflación observada actual"
    if raw_text == "πte = Tasa de inflación esperada" or raw_text.startswith("= Tasa de inflación esperada"):
        return "π_t^e: Tasa de inflación esperada por los agentes económicos"
    if raw_text == "ut = Tasa de desempleo actual" or raw_text.startswith("= Tasa de desempleo actual"):
        return "u_t: Tasa de desempleo observada en el período"
    if raw_text == "un = Tasa natural de desempleo (NAIRU)" or raw_text.startswith("= Tasa natural de desempleo"):
        return "u_n: Tasa natural de desempleo o NAIRU"
    if raw_text.startswith("λ = Parámetro") or raw_text.startswith("= Parámetro que mide"):
        return "λ: Parámetro de sensibilidad de la inflación a la brecha de desempleo"
    if "Si ut>un" in raw_text or raw_text.startswith("Si , se produce una reducción"):
        return "Si u_t > u_n: se generan presiones deflacionarias (la tasa de inflación se reduce)."
    if "Si ut<un" in raw_text or raw_text.startswith("Si , se produce un aumento"):
        return "Si u_t < u_n: se generan presiones inflacionarias (la tasa de inflación aumenta)."
    if "Si ut=un" in raw_text or raw_text.startswith("Si , la tasa de inflación no varía"):
        return "Si u_t = u_n: la inflación se mantiene estable en el nivel esperado (equilibrio NAIRU)."
    
    # Fórmulas de la Sección 6 (Stocks/Flujos y Beveridge)
    if "∆U=s·E-f·U=0" in raw_text or "s·E-f·U=0" in raw_text:
        return "Condición de equilibrio dinámico:  ΔU = s · E - f · U = 0"
    if "u*=ss+f" in raw_text or "u*=" in raw_text:
        return "Tasa de desempleo de estado estacionario:  u* = s / (s + f)"
    if "m=fU,V,A" in raw_text or raw_text == "m=f(U,V,A)":
        return "Función de emparejamiento (matching function):  m = f(U, V, A) = A · U^α · V^(1-α)"
    if raw_text == "m = Número de emparejamientos (contrataciones) por período" or raw_text.startswith("= Número de emparejamientos"):
        return "m: Número de contrataciones o emparejamientos realizados por período"
    if raw_text == "U = Número de desempleados al comienzo del período" or raw_text.startswith("= Número de desempleados"):
        return "U: Stock de personas desempleadas al inicio del período"
    if raw_text == "V = Número de vacantes al comienzo del período" or raw_text.startswith("= Número de vacantes"):
        return "V: Stock de puestos de trabajo vacantes disponibles en la economía"
    if raw_text == "A = Eficiencia del proceso de emparejamiento" or raw_text.startswith("= Eficiencia del proceso"):
        return "A: Parámetro de eficiencia tecnológica e institucional del emparejamiento"

    # Errata histórica y remisión a figuras en Sección 6.4
    if raw_text == "Dos fases distintas identificadas:":
        return "Tres fases macroeconómicas identificadas en la serie histórica (ver Panel B de la Figura 1.9):"
    if raw_text in ["Fase 1 (1994-2007): Estabilidad pre-crisis", "Fase 1 (1980-2007): Estabilidad pre-crisis"]:
        return "Fase 1 (1980–2007: Estabilidad pre-crisis y ancla en Punto A)"
    if raw_text == "Fase 2 (2008-2014): Crisis y desplazamiento":
        return "Fase 2 (2008–2013: Gran Recesión y desplazamiento hacia el Punto B)"
    if raw_text == "Fase 3 (2014-2016): Recuperación única":
        return "Fase 3 (2014–2024: Recuperación económica y récord histórico de vacantes)"

    # Normalización de numeración en epígrafes principales
    if raw_text.strip() == "EQUILIBRIO EN EL MERCADO DE TRABAJO":
        return "3. EQUILIBRIO EN EL MERCADO DE TRABAJO"

    return raw_text

def style_topic_document(input_path: str, output_path: str):
    """Procesa el documento original y genera la versión maquetada con alta fidelidad y figuras."""
    print(f"Leyendo documento original: {input_path}")
    doc_orig = docx.Document(input_path)
    
    # Crear nuevo documento con estilos y geometría
    doc = docx.Document()
    set_document_geometry(doc)
    setup_document_styles(doc)
    
    # Rutas relativas del tema y figuras
    project_root = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
    theme_dir = os.path.dirname(os.path.dirname(os.path.abspath(input_path)))
    fig_dir = os.path.join(theme_dir, "figuras")
    orig_fig_dir = os.path.join(fig_dir, "originales")
    remaster_fig_dir = os.path.join(fig_dir, "remasterizadas")
    os.makedirs(orig_fig_dir, exist_ok=True)
    os.makedirs(remaster_fig_dir, exist_ok=True)
    
    # 1. Extracción automática de figuras originales del .docx
    rid_map = extract_media_from_docx(input_path, orig_fig_dir)
    
    # 2. Generación de figuras remasterizadas mediante psll-figuras-skill
    figuras_skill_dir = os.path.join(project_root, ".skills", "psll-figuras-skill")
    if figuras_skill_dir not in sys.path:
        sys.path.insert(0, figuras_skill_dir)
    try:
        from generate_topic_figures import generate_figures_for_topic
        generate_figures_for_topic(1, project_root)
    except Exception as e:
        print(f"Nota sobre generador de figuras: {e}")
    
    # Rutas de assets oficiales
    cand_emblems = [
        os.path.join(project_root, "Logos y skills", "logo_psll_upo.png"),
        os.path.join(project_root, "Logos y skills", "psll_emblem.png"),
        os.path.join(project_root, "Logos y skills", "psll-presentaciones-skill", "assets", "logo_psll_upo.png"),
        os.path.join(project_root, "Logos y skills", "psll-presentaciones-skill", "assets", "psll_emblem.png"),
    ]
    emblem_path = next((p for p in cand_emblems if os.path.exists(p)), cand_emblems[0])
    upo_logo_path = os.path.join(project_root, "Logos y skills", "psll-presentaciones-skill", "assets", "upo_logo.jpg")
    
    # Cabecera Institucional
    create_institutional_cover(doc, emblem_path, upo_logo_path)
    
    # Identificar título del tema para encabezados dinámicos
    topic_title = "Tema 1: Fundamentos del Mercado Laboral"
    
    # Variables de control
    inserted_figures = set()
    fig_counter = 1
    skip_reflection_block = False
    skip_figures_proposal = False
    inserted_sessions = set()
    collecting_key_concepts = False
    key_concepts = []
    key_concepts_title = "Conceptos Clave del Tema"
    
    print("Procesando y reclasificando párrafos...")
    for idx, p in enumerate(doc_orig.paragraphs):
        raw_text = clean_and_enhance_text(extract_paragraph_full_text(p))
        
        # 0. DETECCIÓN E INSERCIÓN INTELIGENTE DE FIGURAS (ORIGINALES O REMASTERIZADAS)
        blips = p._p.xpath('.//a:blip/@r:embed')
        if blips:
            for r_id in blips:
                fname = rid_map.get(r_id)
                if fname and fname not in inserted_figures:
                    inserted_figures.add(fname)
                    meta = TEMA_1_FIGURES_META.get(fname, {})
                    if meta.get("skip", False):
                        continue
                    remastered_name = meta.get("remastered")
                    caption = meta.get("caption", f"Figura 1.{fig_counter}: Ilustración complementaria del epígrafe.")
                    source = meta.get("source", "Fuente: Elaboración propia para Políticas Sociolaborales (UPO).")
                    
                    # Precedencia: 1) Remasterizada si existe, 2) Original extraída
                    target_fig = None
                    if remastered_name:
                        cand_remaster = os.path.join(remaster_fig_dir, remastered_name)
                        if os.path.exists(cand_remaster):
                            target_fig = cand_remaster
                            
                    if not target_fig:
                        cand_orig = os.path.join(orig_fig_dir, fname)
                        if os.path.exists(cand_orig):
                            target_fig = cand_orig
                            
                    if target_fig:
                        print(f"  -> Insertando figura #{fig_counter} ({fname}): {os.path.basename(target_fig)}")
                        insert_figure(doc, target_fig, caption_text=caption, source_text=source)
                        fig_counter += 1
                        if fname == "image1.png":
                            insert_epa_summary_table(doc)
                        elif fname == "image8.png":
                            # Párrafo explicativo conectando con el doble panel de la Figura 1.8
                            p_exp = doc.add_paragraph(style='Normal')
                            p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            p_exp.paragraph_format.space_before = Pt(4)
                            p_exp.paragraph_format.space_after = Pt(6)
                            p_exp.paragraph_format.line_spacing = 1.15
                            r_exp = p_exp.add_run(
                                "Interpretación analítica de la Figura 1.8: En el Panel A (modelo teórico de Friedman-Phelps), "
                                "un estímulo de demanda reduce coyunturalmente el desempleo a corto plazo desplazando la economía de A a B sobre CP₁; "
                                "sin embargo, la revisión al alza de las expectativas de inflación desplaza la curva a CP₂, retornando "
                                "el desempleo a la NAIRU en C (curva vertical a largo plazo LP). En el Panel B, la serie empírica de España "
                                "(2002–2024) confirma esta dinámica: la fase de burbuja (2002-2007) dio paso a la severa devaluación salarial "
                                "de 2012 (-3,6% con paro al 26%) y, tras la recuperación, al reciente shock inflacionista post-COVID con aumentos de costes salariales superiores al 5%."
                            )
                            r_exp.font.name = FONT_BODY
                            r_exp.font.size = Pt(9.8)
                            r_exp.font.italic = True
                            r_exp.font.color.rgb = RGB_COLORS["ink_green"]
                        elif fname == "image9.png":
                            # Párrafo explicativo conectando con el doble panel de la Figura 1.9
                            p_exp = doc.add_paragraph(style='Normal')
                            p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            p_exp.paragraph_format.space_before = Pt(4)
                            p_exp.paragraph_format.space_after = Pt(6)
                            p_exp.paragraph_format.line_spacing = 1.15
                            r_exp = p_exp.add_run(
                                "Interpretación analítica de la Figura 1.9: El Panel A sintetiza el modelo teórico de emparejamiento (DMP): "
                                "los ciclos económicos generan oscilaciones a lo largo de la curva (expansión vs. recesión), mientras que los "
                                "desajustes estructurales de cualificación (mismatch) desplazan la curva hacia el exterior (E₁ → E₂). "
                                "El Panel B muestra la estimación empírica en España (1980–2024): tras el shock de la Gran Recesión que desplazó "
                                "la curva del Punto A al Punto B (+5 p.p. de paro estructural), la recuperación reciente hasta 2024 ha situado las vacantes "
                                "en máximos históricos (0,73%) mientras el desempleo se sitúa en el 11,3%, confirmando un problema persistente de emparejamiento."
                            )
                            r_exp.font.name = FONT_BODY
                            r_exp.font.size = Pt(9.8)
                            r_exp.font.italic = True
                            r_exp.font.color.rgb = RGB_COLORS["ink_green"]
                            
        if not raw_text:
            continue

        # 0.-1. FILTRADO DE APÉNDICE RESIDUAL DE PROPUESTAS DE FIGURAS
        if raw_text.startswith("2. CURVAS DE OFERTA Y DEMANDA DE TRABAJO") or "GRÁFICOS COMPLEMENTARIOS" in raw_text or "TABLAS RECOMENDABLES" in raw_text or "ELEMENTOS ADICIONALES OPCIONALES" in raw_text:
            skip_figures_proposal = True
            
        if skip_figures_proposal:
            if re.match(r"^Conceptos\s+clave", raw_text, re.IGNORECASE):
                skip_figures_proposal = False
            else:
                continue

        # 0.-2. INSERCIÓN DE CAJAS DE SESIÓN DOCENTE (SESSION BOXES) SEGÚN CRONOGRAMA
        for sess in TEMA_1_SESSION_BOXES:
            if sess["id"] not in inserted_sessions and re.search(sess["pattern"], raw_text, re.IGNORECASE):
                insert_callout_box(
                    doc,
                    callout_type="session",
                    title=sess["title"],
                    text=sess["text"]
                )
                inserted_sessions.add(sess["id"])
                break
            
        # 0.0. RECOLECCIÓN Y VOLCADO DE CONCEPTOS CLAVE EN PÁGINA NUEVA
        if collecting_key_concepts:
            # Si encontramos un nuevo encabezado (capítulo, debate, referencias, etc.), cerramos la caja
            if re.match(r"^\d+\.", raw_text) or raw_text in ["PREGUNTAS PARA EL DEBATE", "REFERENCIAS BIBLIOGRÁFICAS"] or raw_text.startswith("TEMA"):
                if key_concepts:
                    insert_key_concepts_box(doc, title=key_concepts_title, concepts_list=key_concepts, page_break_before=True)
                collecting_key_concepts = False
                key_concepts = []
                # Se continúa la ejecución sin continue para que el párrafo actual se procese normalmente
            else:
                if ":" in raw_text:
                    parts = raw_text.split(":", 1)
                    key_concepts.append((parts[0].strip(), parts[1].strip()))
                elif raw_text:
                    key_concepts.append(("", raw_text.strip()))
                continue

        # Detección del inicio de Conceptos Clave
        if re.match(r"^Conceptos\s+clave", raw_text, re.IGNORECASE):
            collecting_key_concepts = True
            key_concepts_title = raw_text.rstrip(":")
            key_concepts = []
            continue

        # 0.1. PARADA REFLEXIVA AGRUPADA: RANGOS DE LA TASA DE ACTIVIDAD
        if "Reflexión sobre los rangos de la tasa de actividad" in raw_text:
            insert_reflection_box(
                doc,
                title="Límites y Rangos de Variación de la Tasa de Actividad",
                q_and_a_list=[
                    ("¿Puede ser negativa?", "Teóricamente, la tasa de actividad no puede asumir valores negativos, ya que una población activa negativa carece de sentido práctico."),
                    ("¿Es razonable que sea cero?", "Una tasa de actividad del 0% indicaría que ninguna persona en edad de trabajar está buscando empleo ni trabajando, lo cual sería un escenario extremadamente raro y poco realista en cualquier economía moderna."),
                    ("¿Es razonable que sea del 20%?", "Una tasa de actividad del 20% podría darse en países o regiones con una participación limitada en el mercado laboral, quizás debido a restricciones culturales, sociales o a altos niveles de desempleo estructural. No obstante, sigue siendo un valor relativamente bajo en comparación con economías desarrolladas, donde las tasas de actividad suelen oscilar entre el 60% y el 80%.")
                ]
            )
            skip_reflection_block = True
            continue
            
        if skip_reflection_block:
            if "Factores determinantes de la tasa de actividad" in raw_text:
                skip_reflection_block = False
            else:
                continue

        # 0.2. INDICACIÓN DE SESIÓN DOCENTE (CRONOGRAMA DE AULA)
        if "📍" in raw_text or (raw_text.startswith("SESIÓN") and any(k in raw_text for k in ["Semana", "min", "Lunes", "Martes", "Viernes"])):
            clean_session = raw_text.replace("📍", "").strip()
            parts = clean_session.split("—", 1)
            if len(parts) == 2:
                sess_title = parts[0].strip()
                sess_desc = parts[1].strip()
            else:
                sess_title = clean_session
                sess_desc = "Desarrollo presencial de los contenidos, debates y dinámicas activas correspondientes a este bloque según la planificación oficial de la asignatura."
            insert_callout_box(
                doc,
                callout_type="session",
                title=f"Encuadre Docente Presencial · {sess_title}",
                text=sess_desc
            )
            continue

        # 1. TÍTULO PRINCIPAL (TEMA X: ...)
        if re.match(r"^TEMA\s+\d+:", raw_text, re.IGNORECASE):
            topic_title = raw_text
            p_new = doc.add_paragraph()
            p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_new.paragraph_format.space_before = Pt(8)
            p_new.paragraph_format.space_after = Pt(14)
            p_new.paragraph_format.keep_with_next = True
            run = p_new.add_run(raw_text)
            run.font.name = FONT_HEADINGS
            run.font.size = Pt(20.0)
            run.font.bold = True
            run.font.color.rgb = RGB_COLORS["deep_green"]
            continue
            
        # 2. CAPÍTULOS PRINCIPALES (1. FUNDAMENTOS..., 2. EL PIB..., 6. LA CURVA...)
        if re.match(r"^\d+\.\s+[A-ZÁÉÍÓÚÑ\s]{4,}", raw_text) or raw_text in ["EQUILIBRIO EN EL MERCADO DE TRABAJO", "PREGUNTAS PARA EL DEBATE", "REFERENCIAS BIBLIOGRÁFICAS"]:
            p_new = doc.add_paragraph(style='Heading 1')
            p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_new.paragraph_format.space_before = Pt(26)
            p_new.paragraph_format.space_after = Pt(6)
            p_new.paragraph_format.keep_with_next = True
            run = p_new.add_run(raw_text)
            run.font.name = FONT_HEADINGS
            run.font.size = Pt(15.0)
            run.font.bold = True
            run.font.color.rgb = RGB_COLORS["deep_green"]
            continue
            
        # 3. SUBEPÍGRAFES NIVEL 2 (1.1., 1.2., 2.1., 6.1., etc.)
        if re.match(r"^\d+\.\d+\.\s+", raw_text):
            p_new = doc.add_paragraph(style='Heading 2')
            p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_new.paragraph_format.space_before = Pt(18)
            p_new.paragraph_format.space_after = Pt(5)
            p_new.paragraph_format.keep_with_next = True
            run = p_new.add_run(raw_text)
            run.font.name = FONT_HEADINGS
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGB_COLORS["sage"]
            continue
            
        # 4. SUBEPÍGRAFES NIVEL 3 (1.2.1., 1.2.2., 4.1.1., etc.)
        if re.match(r"^\d+\.\d+\.\d+\.\s+", raw_text):
            p_new = doc.add_paragraph(style='Heading 3')
            p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_new.paragraph_format.space_before = Pt(12)
            p_new.paragraph_format.space_after = Pt(3)
            p_new.paragraph_format.keep_with_next = True
            run = p_new.add_run(raw_text)
            run.font.name = FONT_HEADINGS
            run.font.size = Pt(11.0)
            run.font.bold = True
            run.font.color.rgb = RGB_COLORS["deep_green"]
            continue

        # 4.1. FÓRMULAS MATEMÁTICAS DESTACADAS
        is_formula = any(raw_text.startswith(k) for k in [
            "π_t =", "Condición de equilibrio dinámico:", "Tasa de desempleo de estado",
            "Función de emparejamiento", "PIB = VAB", "Tasa de variación interanual:", "PIB per cápita ="
        ])
        if is_formula:
            p_new = doc.add_paragraph()
            p_new.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_new.paragraph_format.space_before = Pt(6)
            p_new.paragraph_format.space_after = Pt(6)
            run = p_new.add_run(raw_text)
            run.font.name = FONT_BODY
            run.font.size = Pt(11.0)
            run.font.bold = True
            run.font.color.rgb = RGB_COLORS["deep_green"]
            continue

        # 5. PÁRRAFOS CONCEPTUALES CON ETIQUETA DESTACADA (PROSA ACADÉMICA CONTINUA, SIN VIÑETAS)
        # Patrón A: Párrafo numerado con concepto clave destacado (ej. "1. Heterogeneidad del factor trabajo: ...")
        m_num_item = re.match(r"^(\d+)\.\s+([A-ZÁÉÍÓÚÑ][^:]{2,55}):\s*(.*)$", raw_text)
        if m_num_item:
            p_new = doc.add_paragraph(style='Normal')
            p_new.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_new.paragraph_format.left_indent = Inches(0)
            p_new.paragraph_format.first_line_indent = Inches(0)
            p_new.paragraph_format.space_before = Pt(2)
            p_new.paragraph_format.space_after = Pt(4.5)
            p_new.paragraph_format.line_spacing = 1.15
            
            run_t = p_new.add_run(f"{m_num_item.group(1)}. {m_num_item.group(2).strip()}: ")
            run_t.font.name = FONT_BODY
            run_t.font.size = Pt(10.5)
            run_t.font.bold = True
            run_t.font.color.rgb = RGB_COLORS["deep_green"]
            
            if m_num_item.group(3).strip():
                run_c = p_new.add_run(m_num_item.group(3).strip())
                run_c.font.name = FONT_BODY
                run_c.font.size = Pt(10.5)
                run_c.font.color.rgb = RGB_COLORS["ink_green"]
            continue

        # Patrón B: Elemento conceptual con etiqueta explicativa antes de dos puntos (ej. "Motor económico: El empleo...")
        m_lbl_item = re.match(r"^([A-ZÁÉÍÓÚÑ][^:]{2,40}):\s+(.+)$", raw_text)
        if m_lbl_item and not raw_text.startswith(("TEMA", "NOTA", "Figura", "Fuente", "Definición", "Interpretación")):
            p_new = doc.add_paragraph(style='Normal')
            p_new.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_new.paragraph_format.left_indent = Inches(0)
            p_new.paragraph_format.first_line_indent = Inches(0)
            p_new.paragraph_format.space_before = Pt(2)
            p_new.paragraph_format.space_after = Pt(4.5)
            p_new.paragraph_format.line_spacing = 1.15
            
            run_t = p_new.add_run(f"{m_lbl_item.group(1).strip()}: ")
            run_t.font.name = FONT_BODY
            run_t.font.size = Pt(10.5)
            run_t.font.bold = True
            run_t.font.color.rgb = RGB_COLORS["deep_green"]
            
            run_c = p_new.add_run(m_lbl_item.group(2).strip())
            run_c.font.name = FONT_BODY
            run_c.font.size = Pt(10.5)
            run_c.font.color.rgb = RGB_COLORS["ink_green"]
            continue

        # Patrón C: Listas de requisitos técnicos estrictos (única excepción tasada para viñetas •)
        is_strict_bullet = raw_text.startswith((
            "Estar sin empleo", "Haber tomado medidas concretas", "Estar disponible para trabajar"
        )) or (
            bool(re.match(r"^(\-|\•|\*)\s*(.+)$", raw_text)) and len(raw_text.split()) <= 25
        )
        if is_strict_bullet:
            clean_item = re.sub(r"^(\-|\•|\*)\s*", "", raw_text).strip()
            p_new = doc.add_paragraph(style='Normal')
            p_new.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_new.paragraph_format.left_indent = Inches(0.35)
            p_new.paragraph_format.first_line_indent = Inches(-0.18)
            p_new.paragraph_format.space_before = Pt(1)
            p_new.paragraph_format.space_after = Pt(3.0)
            p_new.paragraph_format.line_spacing = 1.15
            
            run_b = p_new.add_run("•  ")
            run_b.font.name = FONT_BODY
            run_b.font.bold = True
            run_b.font.color.rgb = RGB_COLORS["sage"]
            
            run_c = p_new.add_run(clean_item)
            run_c.font.name = FONT_BODY
            run_c.font.size = Pt(10.5)
            run_c.font.color.rgb = RGB_COLORS["ink_green"]
            continue

        # 6. PÁRRAFO NORMAL (CUERPO DE TEXTO JUSTIFICADO)
        p_new = doc.add_paragraph(style='Normal')
        p_new.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_new.paragraph_format.left_indent = Inches(0)
        p_new.paragraph_format.first_line_indent = Inches(0)
        p_new.paragraph_format.space_after = Pt(4.5)
        p_new.paragraph_format.line_spacing = 1.15
        
        # Preservar negritas básicas si el primer run original era negrita o títulos internos cortos
        first_bold = False
        is_short_heading = len(raw_text) < 55 and not raw_text.endswith((".", ":", ";"))
        if (p.runs and p.runs[0].bold) or is_short_heading:
            first_bold = True
            
        # Si es un título o subtítulo corto interno, forzar alineación a la izquierda
        if is_short_heading:
            p_new.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_new.paragraph_format.space_before = Pt(8)
            p_new.paragraph_format.space_after = Pt(2)
            p_new.paragraph_format.keep_with_next = True
            
        # Copiar texto
        run = p_new.add_run(raw_text)
        run.font.name = FONT_BODY
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGB_COLORS["ink_green"]
        if first_bold:
            run.font.bold = True
            if is_short_heading:
                run.font.color.rgb = RGB_COLORS["deep_green"]



    if collecting_key_concepts and key_concepts:
        insert_key_concepts_box(doc, title=key_concepts_title, concepts_list=key_concepts, page_break_before=True)

    # Agregar encabezados y pies de página
    add_header_and_footer(doc, topic_title)
    
    # Guardar documento con gestión de archivo abierto en Word
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    try:
        doc.save(output_path)
        print(f"\n✅ Documento maquetado generado con éxito en: {output_path}")
    except PermissionError:
        alt_output = output_path.replace(".docx", "_v2.docx")
        doc.save(alt_output)
        print(f"\n⚠️ El archivo '{output_path}' está abierto en Microsoft Word.")
        print(f"✅ Se ha guardado la versión actualizada en: {alt_output}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Aplica la maquetación corporativa oficial de PSLL a un archivo .docx.")
    cand_input = r"Temas EB\Tema 1\Apuntes\Tema 1 2526.docx"
    if not os.path.exists(cand_input):
        cand_input = r"Temas EB\Tema 1\Apuntes\Tema 1 2627.docx"
    parser.add_argument("--input", default=cand_input, help="Ruta del documento de entrada")
    parser.add_argument("--output", default=r"Temas EB\Tema 1\Apuntes\Tema 1 2627_maquetado.docx", help="Ruta del documento de salida")
    args = parser.parse_args()
    
    style_topic_document(args.input, args.output)
